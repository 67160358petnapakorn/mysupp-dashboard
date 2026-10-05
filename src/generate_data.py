"""สร้างชุดข้อมูลจำลองของ MYSUPP (Star Schema) -> data.xlsx และ src/embed.json
ข้อมูลทั้งหมดเป็นข้อมูลจำลองเพื่อการเรียนการสอน ไม่ใช่ตัวเลขจริงของธุรกิจ
รัน (จากโฟลเดอร์หลักของ repo): python src/generate_data.py แล้วรัน python src/build.py เพื่อสร้าง index.html"""
import random, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
random.seed(69)
TH=['ม.ค.','ก.พ.','มี.ค.','เม.ย.','พ.ค.','มิ.ย.','ก.ค.','ส.ค.','ก.ย.','ต.ค.','พ.ย.','ธ.ค.']
months=[(2025,10),(2025,11),(2025,12)]+[(2026,m) for m in range(1,10)]
NM=len(months)
dkey=lambda y,m:y*10000+m*100+1
labels=[f"{TH[m-1]} {(y+543)%100}" for y,m in months]
SEG=[(1,'วัยทำงาน 25–45 ปี','สุขภาพเชิงป้องกัน ไม่มีเวลาพบแพทย์/นักโภชนาการบ่อยๆ',.52,780),
     (2,'ออกกำลังกาย/ฟิตเนส','ต้องการอาหารเสริมตรงเป้าหมาย เพิ่มกล้ามเนื้อ ฟื้นฟูร่างกาย',.30,920),
     (3,'ผู้สูงอายุ/ผู้ดูแล','ดูแลสุขภาพเชิงป้องกันแบบเฉพาะบุคคล',.18,700)]
CH=[(1,'Social Media','Social Commerce',0.0,45,.30,520),(2,'App / Website MYSUPP','Owned',0.025,45,.27,380),
    (3,'คลินิก/โรงพยาบาลพันธมิตร','Partner',0.10,45,.20,260),(4,'E-commerce (Shopee/Lazada)','Marketplace',0.06,45,.15,610),
    (5,'องค์กร (B2B)','B2B',0.0,20,.08,900)]
new_tot=[80,110,150,190,240,290,340,390,440,490,530,570]
churn=[.18,.165,.15,.135,.12,.11,.10,.09,.085,.08,.075,.07]
adh=[58,60,63,66,69,72,74,76,78,80,81,82]
rows=[];mkt=[];fun=[]
start={(s[0],c[0]):0 for s in SEG for c in CH}
for t,(y,m) in enumerate(months):
    cn={}
    for c in CH:
        n_c=0
        for s in SEG:
            if c[0]==5 and t<3: new=0
            else: new=max(0,round(new_tot[t]*s[3]*c[5]*random.uniform(.9,1.1)))
            a0=start[(s[0],c[0])]; ch=round(a0*churn[t]*random.uniform(.92,1.08)); a1=a0-ch+new
            start[(s[0],c[0])]=a1; n_c+=new
            orders=round(a1*random.uniform(.95,.99)); price=s[4]*random.uniform(.96,1.04)
            rs=round(orders*price); rsub=a1*149; rcon=round(a1*0.10*1200*random.uniform(.9,1.1))
            rb2b=orders*350 if c[0]==5 else 0
            cogs=round(rs*random.uniform(.38,.42)); fees=round((rs+rsub)*c[3]); ship=orders*c[4]
            ad=round(a1*max(0,min(1,random.gauss(adh[t],3)/100)))
            rows.append([len(rows)+1,dkey(y,m),s[0],c[0],new,a1,ch,ad,orders,rs,rsub,rcon,rb2b,cogs,fees,ship])
        cac=c[6]*(1+0.03*t if c[0] in (1,4) else 1+0.01*t)*random.uniform(.95,1.05)
        ad_sp=round(n_c*cac*0.8); kol=round(n_c*cac*0.2) if c[0]==1 else 0
        if c[0]!=1: ad_sp=round(n_c*cac)
        imp=round((ad_sp+kol)*random.uniform(55,70)); clk=round(imp*random.uniform(.012,.02))
        mkt.append([len(mkt)+1,dkey(y,m),c[0],ad_sp,kol,imp,clk,n_c])
        up_rate=min(.8,(.40+.012*t)*(1.7 if c[0]==3 else 1)); up=round(n_c/.62) if n_c else 0
        q=round(up/up_rate) if n_c else 0; vis=round(q/.18) if n_c else 0
        fun.append([len(fun)+1,dkey(y,m),c[0],vis,q,up,n_c,round(n_c*random.uniform(.68,.74)),round(n_c*random.uniform(.16,.22))])
opex=[[dkey(y,m),210000+12000*t,70000+5000*t,45000+2000*t,25000+3000*t,30000,20000+1000*t] for t,(y,m) in enumerate(months)]
MK=[('แมกนีเซียม',72),('วิตามินดี',68),('ธาตุเหล็ก',61),('โอเมก้า-3 / ไขมันในเลือด',57)]
health=[]
for mi,(mn,rate) in enumerate(MK,1):
    for s in SEG:
        tested=random.randint(140,210); imp=round(tested*min(.95,max(.3,rate/100+random.uniform(-.04,.04))))
        health.append([len(health)+1,mi,mn,s[0],tested,imp])
# ---------- workbook
wb=Workbook(); H=Font(name='Arial',bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='123B5D'); F=Font(name='Arial')
def sheet(name,head,data,widths=None):
    ws=wb.create_sheet(name); ws.append(head)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(horizontal='center')
    for r in data: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=F
    for i,col in enumerate(ws.columns): ws.column_dimensions[col[0].column_letter].width=(widths[i] if widths else 16)
    ws.freeze_panes='A2'; return ws
wb.remove(wb.active)
rd=wb.create_sheet('README')
for line in ['MYSUPP — ชุดข้อมูลจำลอง (Star Schema)','รายวิชา Business Idea Creation · จัดทำโดย เพชรนภากร พลนิกร 67160358',
 f'ข้อมูลทั้งหมดเป็นข้อมูลจำลอง ช่วง {labels[0]} – {labels[-1]} ({NM} เดือน) สกุลเงินบาท ไม่ใช่ตัวเลขจริงของธุรกิจ','',
 'ชีต Summary: ตัวเลขสรุปคำนวณด้วยสูตรจากชีต fact_*','fact_monthly: 1 แถว = 1 เดือน × 1 กลุ่มลูกค้า × 1 ช่องทาง (สมาชิก ยอดขาย ต้นทุน)',
 'fact_marketing_monthly: ค่าโฆษณา/KOL และสมาชิกใหม่ ต่อเดือนต่อช่องทาง','fact_funnel_monthly: ผู้ใช้แต่ละขั้นของ Funnel ต่อเดือนต่อช่องทาง',
 'fact_opex_monthly: ค่าใช้จ่ายคงที่รายเดือน (ทีมงาน AI/Cloud แอป ลูกค้าสัมพันธ์ ลิขสิทธิ์ ความปลอดภัย)','fact_health_outcome: ผู้ใช้ที่ค่าตรวจเลือดดีขึ้นหลังใช้ครบ 3 เดือน แยกตัวชี้วัด×กลุ่มลูกค้า',
 'dim_*: ตารางมิติ (เดือน กลุ่มลูกค้า ช่องทาง ตัวชี้วัดสุขภาพ)','','สมมติฐานสำคัญ: ราคาเฉลี่ย/ออเดอร์ 700–920 บาท · ค่าสมาชิก 149 บาท/คน/เดือน · ที่ปรึกษา Premium 10% ของสมาชิก ×1,200 บาท · ต้นทุนสินค้า ≈ 40% ของยอดขายอาหารเสริม',
 'สร้างด้วย src/generate_data.py (random seed คงที่) — แก้ตัวเลขแล้วรันใหม่ได้']: rd.append([line])
rd['A1'].font=Font(name='Arial',bold=True,size=14); rd.column_dimensions['A'].width=110
sheet('fact_monthly',['row_key','month_key','segment_key','channel_key','new_members','active_members_end','churned_members','adherent_members','orders','supplement_revenue','subscription_revenue','consult_revenue','b2b_revenue','cogs','platform_fee','shipping_cost'],rows)
sheet('fact_marketing_monthly',['row_key','month_key','channel_key','ad_spend','kol_cost','impressions','clicks','new_members'],mkt)
sheet('fact_funnel_monthly',['row_key','month_key','channel_key','visitors','questionnaire_done','blood_result_uploaded','subscribed','renewed_month3','referred_friend'],fun)
sheet('fact_opex_monthly',['month_key','team_cost','ai_cloud_cost','app_dev_cost','customer_service_cost','software_license_cost','cybersecurity_cost'],opex)
sheet('fact_health_outcome',['row_key','marker_key','marker_name','segment_key','users_tested','users_improved'],health,[10,12,28,13,14,14])
sheet('dim_month',['month_key','year','month_num','label_th'],[[dkey(y,m),y,m,labels[i]] for i,(y,m) in enumerate(months)])
sheet('dim_segment',['segment_key','segment_name','description_th'],[[s[0],s[1],s[2]] for s in SEG],[13,26,70])
sheet('dim_channel',['channel_key','channel_name','channel_group','fee_rate','shipping_cost_per_order'],[[c[0],c[1],c[2],c[3],c[4]] for c in CH],[13,30,18,10,22])
n=len(rows)+1; nm=len(mkt)+1; no=len(opex)+1
rng=lambda col,sh='fact_monthly',last=n:f"{sh}!${col}$2:${col}${last}"
sm=wb.create_sheet('Summary',1); sm.append(['สรุปตัวเลขหลัก (คำนวณด้วยสูตรจากชีต fact_*)']); sm.append(['ตัวชี้วัด','ค่า'])
last=dkey(*months[-1])
sm_rows=[('ยอดขายสุทธิรวม (บาท)',f"=SUM({rng('J')})+SUM({rng('K')})+SUM({rng('L')})+SUM({rng('M')})"),
 ('ต้นทุนสินค้า (บาท)',f"=SUM({rng('N')})"),('กำไรขั้นต้น (%)','=1-B4/B3'),
 ('ค่าธรรมเนียม+ค่าจัดส่ง (บาท)',f"=SUM({rng('O')})+SUM({rng('P')})"),
 ('ค่าการตลาด (บาท)',f"=SUM({rng('D','fact_marketing_monthly',nm)})+SUM({rng('E','fact_marketing_monthly',nm)})"),
 ('ค่าใช้จ่ายคงที่ (บาท)',f"=SUM(fact_opex_monthly!$B$2:$G${no})"),
 ('กำไรสุทธิ (บาท)','=B3-B4-B6-B7-B8'),
 ('สมาชิกที่ใช้งานอยู่เดือนล่าสุด (คน)',f"=SUMIFS({rng('F')},{rng('B')},{last})"),
 ('สมาชิกใหม่ทั้งหมด (คน)',f"=SUM({rng('E')})"),
 ('CAC เฉลี่ย (บาท/คน)','=B7/B11')]
for r in sm_rows: sm.append(list(r))
sm['A1'].font=Font(name='Arial',bold=True,size=13)
for c in sm[2]: c.font=H; c.fill=HF
for r in sm.iter_rows(min_row=3):
    for c in r: c.font=F
sm.column_dimensions['A'].width=44; sm.column_dimensions['B'].width=20
for r in (3,4,6,7,8,9,10,11,12,13): sm[f'B{r}'].number_format='#,##0'
sm['B5'].number_format='0.0%'
wb.save('data.xlsx')
# ---------- embed json (เดือนใช้เป็นลำดับ 0..11)
IDX={dkey(y,m):i for i,(y,m) in enumerate(months)}
json.dump({'labels':labels,'seg':[[s[0],s[1],s[2]] for s in SEG],'ch':[[c[0],c[1]] for c in CH],
 'rows':[[IDX[r[1]],r[2],r[3]]+r[4:] for r in rows],'mkt':[[IDX[m[1]],m[2]]+m[3:] for m in mkt],
 'fun':[[IDX[f[1]],f[2]]+f[3:] for f in fun],'opex':[o[1:] for o in opex],'health':[[h[2],h[3],h[4],h[5]] for h in health]},
 open('src/embed.json','w'),ensure_ascii=False,separators=(',',':'))
